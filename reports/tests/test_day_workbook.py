"""The daily report's Excel download: every table of the page in one file, real numbers, and the data label."""
import json

from openpyxl import load_workbook

from reports import day_workbook, pulse

EXAMPLES = pulse.ROOT / "docs" / "contracts" / "examples"
CLEAN = json.loads((EXAMPLES / "pulse.sample.json").read_text(encoding="utf-8"))
MISSING = json.loads((EXAMPLES / "pulse.sample.missing.json").read_text(encoding="utf-8"))


def sheet(ws):
    return [list(row) for row in ws.iter_rows(values_only=True)]


def test_one_sheet_per_table_each_labelled_synthetic(tmp_path):
    assert day_workbook.write(tmp_path / "day.xlsx", CLEAN) is True
    wb = load_workbook(tmp_path / "day.xlsx")
    assert wb.sheetnames == ["Summary", "By Marketplace", "Data Quality", "Definitions"]
    for ws in wb:
        assert ws["A1"].value.startswith("Synthetic sample data"), ws.title
    rows = sheet(wb["By Marketplace"])
    assert rows[2][:3] == ["Marketplace", "Status", "Revenue"]
    by_name = {r[0]: r for r in rows[3:]}
    assert (by_name["ShopGoodwill"][2], by_name["Amazon"][2], by_name["eBay"][2]) == (209.0, 47.98, 18.5)
    assert by_name["Enterprise Total"][2] == 275.48          # dollars as numbers, so Excel can sum them
    assert by_name["Amazon"][7] == "order" and by_name["eBay"][7] == "buyer"
    assert dict(r[:2] for r in sheet(wb["Summary"])[3:])["Enterprise revenue"] == 275.48


def test_a_marketplace_with_no_data_is_blank_never_zero(tmp_path):
    day_workbook.write(tmp_path / "day.xlsx", MISSING)
    wb = load_workbook(tmp_path / "day.xlsx")
    ebay = next(r for r in sheet(wb["By Marketplace"]) if r[0] == "eBay")
    assert ebay[1] == "missing" and ebay[2:] == [None] * 8
    assert dict(r[:2] for r in sheet(wb["Summary"])[3:])["Totals"] == "Partial: no data for eBay"


def test_the_nights_transactions_and_flagged_rows_come_from_the_engines_folder(tmp_path):
    engine = tmp_path / "engine"
    engine.mkdir()
    (engine / "transactions.csv").write_text(
        "txn_id,marketplace,business_date,gross_cents,fee_cents\n"
        "a:1,amazon,2026-10-02,2999,450\na:2,amazon,2026-10-01,1000,150\ne:1,ebay,2026-10-02,-500,\n", encoding="utf-8")
    (engine / "warnings.json").write_text(json.dumps(
        [{"source_file": "ebay.csv", "source_row": 4, "kind": "duplicate", "reason": "same transaction as row 2"}]), encoding="utf-8")
    day_workbook.write(tmp_path / "day.xlsx", CLEAN, engine)
    wb = load_workbook(tmp_path / "day.xlsx")
    assert wb.sheetnames[-2:] == ["Transactions", "Flagged Rows"]
    rows = sheet(wb["Transactions"])
    assert rows[2] == ["Txn id", "Marketplace", "Business date", "Gross ($)", "Fee ($)"]
    assert rows[3:] == [["a:1", "amazon", "2026-10-02", 29.99, 4.5], ["e:1", "ebay", "2026-10-02", -5.0, None]]  # that day only
    assert sheet(wb["Flagged Rows"])[3] == ["ebay.csv", 4, "duplicate", "same transaction as row 2"]


def test_the_page_offers_the_workbook_and_keeps_explanations_for_the_notes(tmp_path):
    paths = pulse.render(CLEAN, tmp_path)
    assert paths[-1].name == "2026-10-02.xlsx" and paths[-1].exists()
    html = paths[0].read_text(encoding="utf-8")
    assert '<a class="btn dl" href="2026-10-02.xlsx" download>Download Full Day (Excel)</a>' in html
    assert "<h1>Nightly Pulse: Friday, October 2, 2026</h1>" in html
    table, bottom = html.split('<details class="notes" id="notes">')
    assert "counted by order" not in table.split("<tbody>")[1]        # an asterisk in the table...
    assert table.count('class="ast"') >= 4
    assert "Amazon: counted by order, because the export gives no unique buyer id" in bottom   # ...the reason at the bottom
    assert "Definitions" in bottom and "Data Quality" in bottom
