"""The full day as one Excel file: reports/pulse/<date>.xlsx, behind the daily report's download button.

One sheet per table of the nightly pulse page, with real numbers (dollars, not text) so they sum and chart
in Excel: Summary, By Marketplace, Data Quality and Definitions, all from the pulse JSON. When the engine's
output folder for the night is given, two more: Transactions (the cleaned rows the figures were computed
from) and Flagged Rows (what the engine removed or could not read, and why). No arithmetic here beyond
cents to dollars. Every sheet says the data is synthetic, as the pages do.

Needs openpyxl (engine/requirements.txt). Without it nothing is written and the page shows no button.
"""
import csv
import json
from datetime import date
from pathlib import Path

ORDER = ["shopgoodwill", "amazon", "ebay", "other"]
LABELS = {"shopgoodwill": "ShopGoodwill", "amazon": "Amazon", "ebay": "eBay", "other": "Other"}
MONEY = '"$"#,##0.00;[Red]-"$"#,##0.00'
DATA_LABEL = "Synthetic sample data: no real Goodwill file has been read."


def dollars(cents):
    return None if cents is None else cents / 100


def _sheet(wb, title, header, rows, money_cols=(), widths=None, first=False):
    """A sheet with the data label on top, a bold header row and its rows; returns the worksheet."""
    from openpyxl.styles import Font
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append([DATA_LABEL])
    ws["A1"].font = Font(italic=True, color="5C5859")
    ws.append([])
    ws.append(header)
    for cell in ws[3]:
        cell.font = Font(bold=True, color="0054A4")
    for row in rows:
        ws.append(row)
        for col in money_cols:
            ws.cell(row=ws.max_row, column=col).number_format = MONEY
    ws.freeze_panes = "A4"
    for i, width in enumerate(widths or [], start=1):
        ws.column_dimensions[ws.cell(row=3, column=i).column_letter].width = width
    return ws


def write(path, pulse, detail_dir=None):
    """Write the workbook for one pulse. Returns True when written, False when openpyxl is not installed."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except ImportError:
        return False
    markets, ent, day = pulse["marketplaces"], pulse["enterprise"], pulse["business_date"]
    delta = ent.get("delta") or {}
    excluded = [x["marketplace"] if isinstance(x, dict) else x for x in ent.get("excluded") or []]
    prior = pulse.get("prior_date") or delta.get("prior_date")
    wb = Workbook()

    _sheet(wb, "Summary", ["Figure", "Value"], [
        ["Business date", date.fromisoformat(day)],
        ["Compared with", date.fromisoformat(prior) if prior else "no prior day"],
        ["Enterprise revenue", dollars(ent.get("revenue_cents"))],
        ["Change vs prior day", dollars(delta.get("revenue_cents"))],
        ["Change vs prior day, %", delta.get("revenue_pct")],
        ["Orders", ent.get("orders")],
        ["Customers", ent.get("customers")],
        ["Refunds", dollars(ent.get("refunds_cents"))],
        ["Marketplace fees", dollars(ent.get("fees_cents"))],
        ["Totals", "Partial: no data for " + ", ".join(LABELS.get(k, k) for k in excluded) if excluded else "Complete"],
    ], widths=[28, 44], first=True)
    ws = wb["Summary"]
    for row in (6, 7, 11, 12):  # the dollar figures (rows are offset by the label and header rows)
        ws.cell(row=row, column=2).number_format = MONEY
    for row in (4, 5):
        ws.cell(row=row, column=2).number_format = "mm/dd/yyyy"
        ws.cell(row=row, column=2).alignment = Alignment(horizontal="left")

    rows = []
    for key in ORDER:
        m = markets.get(key, {"status": "missing"})
        name, ok = m.get("label") or LABELS.get(key, key), m.get("status") == "ok"
        d = (m.get("delta") or {}) if ok else {}
        rows.append([name, m.get("status", "missing")] + ([
            dollars(m["revenue_cents"]), dollars(d.get("revenue_cents")), d.get("revenue_pct"), m["orders"],
            m["customers"], m.get("customer_basis") or "", dollars(m["refunds_cents"]), dollars(m["fees_cents"])]
            if ok else [None] * 8))
    rows.append(["Enterprise Total", "partial" if excluded else "complete", dollars(ent.get("revenue_cents")),
                 dollars(delta.get("revenue_cents")), delta.get("revenue_pct"), ent.get("orders"), ent.get("customers"),
                 "", dollars(ent.get("refunds_cents")), dollars(ent.get("fees_cents"))])
    ws = _sheet(wb, "By Marketplace",
                ["Marketplace", "Status", "Revenue", "Change vs Prior Day", "Change %", "Orders", "Customers",
                 "Customer Basis", "Refunds", "Fees"], rows, money_cols=(3, 4, 9, 10),
                widths=[20, 16, 14, 20, 10, 9, 11, 15, 12, 12])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    dq = pulse.get("data_quality") or {}
    _sheet(wb, "Data Quality", ["Check", "Rows"],
           [["Issues handled automatically", dq.get("warnings_total", 0)]]
           + [[kind.replace("_", " ").capitalize(), n] for kind, n in sorted((dq.get("by_kind") or {}).items())]
           + [["Rows that could not be read (left out)", dq.get("rows_rejected", 0)]], widths=[40, 10])

    _sheet(wb, "Definitions", ["Term", "Definition"],
           [[k.replace("_", " ").capitalize(), v] for k, v in (pulse.get("definitions") or {}).items()], widths=[24, 110])

    folder = Path(detail_dir) if detail_dir else None
    if folder and (folder / "transactions.csv").exists():
        with open(folder / "transactions.csv", encoding="utf-8", newline="") as f:
            reader = csv.DictReader(f)
            cols = reader.fieldnames or []
            found = [r for r in reader if r.get("business_date") == day]
        cents = [c for c in cols if c.endswith("_cents")]
        header = [c[:-6].replace("_", " ").capitalize() + " ($)" if c in cents else c.replace("_", " ").capitalize() for c in cols]
        body = [[(int(r[c]) / 100 if r[c] not in ("", None) else None) if c in cents else r[c] for c in cols] for r in found]
        _sheet(wb, "Transactions", header, body, money_cols=[cols.index(c) + 1 for c in cents],
               widths=[max(12, min(40, len(h) + 6)) for h in header])
    if folder and (folder / "warnings.json").exists():
        flagged = json.loads((folder / "warnings.json").read_text(encoding="utf-8"))
        if isinstance(flagged, list) and flagged:
            keys = list(flagged[0])
            _sheet(wb, "Flagged Rows", [k.replace("_", " ").capitalize() for k in keys],
                   [[w.get(k) if isinstance(w.get(k), (str, int, float, type(None))) else json.dumps(w.get(k)) for k in keys]
                    for w in flagged], widths=[max(14, min(70, len(k) + 30)) for k in keys])

    wb.save(path)
    return True
