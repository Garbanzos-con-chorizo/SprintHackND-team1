"""Reads any CSV or XLSX export into a Table: a header plus rows keyed by header name.

Parsers never open files; they get a Table. That keeps format handling (encodings,
delimiters, sheets) in one place so adding a source never touches it.
"""
import csv
import re
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED_SUFFIXES = {".csv", ".tsv", ".txt", ".xlsx"}

# Cell text that exports use to mean "nothing here" (eBay writes '--').
EMPTY_MARKERS = {"", "--"}


class UnreadableFile(Exception):
    """The file could not be read as a table at all."""


def norm(name: object) -> str:
    """Normalize a column name for matching: 'Order ID ' / 'order_id' / 'ORDER-ID' -> 'order id'."""
    return re.sub(r"[^a-z0-9]+", " ", str(name).lower()).strip()


@dataclass
class Table:
    name: str                     # file name, e.g. 'ebay_2026-10-02.csv'
    header: list[str]             # column names as written in the file
    rows: list[tuple[int, dict]]  # (1-based data row number, {header name: cell text})
    norm_header: set[str] = field(init=False)

    def __post_init__(self):
        self.norm_header = {norm(h) for h in self.header}

    def get(self, row: dict, *aliases: str) -> str:
        """First non-empty cell among the aliased column names (compared normalized), else ''."""
        by_norm = {norm(h): h for h in self.header}
        for alias in aliases:
            col = by_norm.get(norm(alias))
            if col is not None:
                text = str(row.get(col, "")).strip()
                if text not in EMPTY_MARKERS:
                    return text
        return ""


def _is_blank(values) -> bool:
    return all(v is None or str(v).strip() == "" for v in values)


def _build(name: str, raw_rows: list[list]) -> Table:
    raw_rows = [r for r in raw_rows]
    # The header is the first row with two or more filled cells; anything above it (a title,
    # a report date) is skipped. A file with only one column falls back to its first non-blank row.
    filled = [sum(1 for v in r if v is not None and str(v).strip() != "") for r in raw_rows]
    start = next((i for i, n in enumerate(filled) if n >= 2), None)
    if start is None:
        start = next((i for i, n in enumerate(filled) if n >= 1), None)
    if start is None:
        raise UnreadableFile("file is empty")
    header = [str(h).strip() if h is not None else "" for h in raw_rows[start]]
    rows = []
    for values in raw_rows[start + 1:]:
        if _is_blank(values):
            continue  # blank rows are ignored without a warning, and not numbered (contract)
        cells = {h: ("" if i >= len(values) or values[i] is None else values[i]) for i, h in enumerate(header) if h}
        rows.append((len(rows) + 1, cells))
    return Table(name=name, header=[h for h in header if h], rows=rows)


def read_table(path: Path) -> Table:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise UnreadableFile(f"unsupported file type {suffix or '(none)'}")
    try:
        if suffix == ".xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(path, read_only=True, data_only=True)
            try:
                raw = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
            finally:
                wb.close()
        else:
            text = path.read_text(encoding="utf-8-sig", errors="replace")
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            raw = list(csv.reader(text.splitlines(), dialect))
    except UnreadableFile:
        raise
    except Exception as exc:  # corrupt xlsx, locked file, etc: warn, never crash the run
        raise UnreadableFile(f"{type(exc).__name__}: {exc}") from exc
    return _build(path.name, raw)
